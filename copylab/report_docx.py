"""Optional Word exporter; run with a Python that provides python-docx."""
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from copylab.report import dollars, timestamp


def clean(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(value))


def table(document, headings, rows, widths):
    grid = document.add_table(rows=1, cols=len(headings))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for column, width in zip(grid.columns, widths):
        column.width = Inches(width)
    properties = grid._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        border = OxmlElement('w:'+side)
        for key, value in (('val','single'),('sz','4'),('color','D9D9D9')):
            border.set(qn('w:'+key), value)
        borders.append(border)
    properties.append(borders)
    for i, heading in enumerate(headings):
        grid.rows[0].cells[i].text = heading
    repeat = OxmlElement('w:tblHeader')
    grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = grid.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = clean(value)
    for row_index, row in enumerate(grid.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                element = OxmlElement('w:'+side)
                element.set(qn('w:w'), '85')
                element.set(qn('w:type'), 'dxa')
                margins.append(element)
            tcpr.append(margins)
            if row_index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EAEAEA')
                tcpr.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.bold = row_index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return grid


def build(report, output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = document.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (('Title',22),('Heading 1',14),('Heading 2',11)):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.bold = name != 'Title'
        style.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(7)
    document.core_properties.title = 'Paper Trading Trade and Profit Report'
    document.core_properties.author = 'CopyLab'
    document.add_paragraph('Paper Trading Trade and Profit Report', 'Title')
    document.add_paragraph('Updated '+report['generated'])
    account = report['account']
    if report['trades']:
        document.add_paragraph('The paper account has recorded %d buys, %d copied sells, and %d settled positions. Total profit or loss is %s. This report lists every recorded trade and separates realized gains from estimated open gains.' %
                               (account['fills'],account['sales'],account['settled_positions'],dollars(account['total_pnl'])))
    else:
        document.add_paragraph('No qualifying copied trades have been recorded yet. The simulated account remains at $40 with $0 profit or loss. This report will update as the running bot records trades.')
    document.add_paragraph('This is a simulated account. No real money is traded. Synthetic test scenarios are excluded.')
    document.add_paragraph('Account summary','Heading 1')
    table(document, ['Measure','Amount'], [
        ('Starting balance','$40.00'),('Available cash',dollars(account['cash'])),
        ('Estimated account equity',dollars(account['equity'])),('Total profit or loss',dollars(account['total_pnl'])),
        ('Realized profit or loss',dollars(account['realized_pnl'])),
        ('Estimated open profit or loss',dollars(account['unrealized_pnl'])),
        ('Entry and exit fee reserves',dollars(report['fee_reserve_total']))], [4.7,2.3])
    document.add_paragraph('Fees use fixed simulation reserves of 2 percent on entry and exit, not verified exchange charges. Open values use available bids after the exit reserve. Missing or stale quotes make total profit unknown. Realized gains come from copied sales and confirmed paper settlements. Buys have no realized gain on entry; their profit is recorded on exit or settlement.')
    if account['positions']:
        document.add_paragraph('Open positions','Heading 1')
        for position in account['positions']:
            document.add_paragraph(clean(position['title'])+' '+clean(position['outcome']), 'Heading 2')
            table(document, ['Shares','Remaining cost','Estimated value','Open profit or loss'],
                  [('%.6f'%position['shares'],dollars(position['cost']),dollars(position['value']),dollars(position['pnl']))],
                  [1.3,1.7,1.7,2.3])
    document.add_paragraph('Complete trade history','Heading 1')
    if not report['trades']:
        document.add_paragraph('There are no copied buys, sells, or settlements to list yet. The CSV export will contain every event as activity arrives.')
    for event in report['trades']:
        document.add_paragraph('%d %s %s' % (event['number'],event['type'],clean(event['market'])), 'Heading 2')
        document.add_paragraph(event['time']+' | '+event['time_basis']+' | Outcome '+clean(event['outcome']))
        table(document, ['Shares','Price per share','Fee reserve','Cash change','Realized gain or loss'],
              [('%.6f'%event['shares'],'$%.4f'%event['price'],dollars(event['fee_reserve']),
                dollars(event['cash_change']), 'On exit' if event['realized_pnl'] is None else dollars(event['realized_pnl']))],
              [1.2,1.4,1.2,1.4,1.8])
        document.add_paragraph('Allocated entry cost '+dollars(event['cost_basis'])+'. Source wallet '+clean(event['wallet'])+'.')
        if event['source_transaction']:
            paragraph = document.add_paragraph('Source transaction '+clean(event['source_transaction']))
            for run in paragraph.runs:
                run.font.size = Pt(8)
    if account['scans']:
        document.add_paragraph('Monitoring progress','Heading 1')
        for scan in account['scans']:
            state = scan['error'] or ('Polling normally' if report['generated_timestamp']-scan['checked'] <= 120 else 'Last scan is stale')
            paragraph = document.add_paragraph(clean(scan['wallet']))
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                run.font.size = Pt(8)
            document.add_paragraph('Last scan '+timestamp(scan['checked'])+'. '+clean(state)+'.')
    output = Path(output)
    temporary = output.with_name(output.stem+'.tmp.docx')
    for root in (document._element, document.styles.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    document.save(temporary)
    os.replace(temporary,output)


if __name__ == '__main__':
    build(json.loads(Path(sys.argv[1]).read_text()), sys.argv[2])
